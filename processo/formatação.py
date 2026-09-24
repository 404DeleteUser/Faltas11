import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ============================================================
# 1. DICIONÁRIOS DE COLUNAS (ORDEM E NOMES)
# ============================================================
DICIONARIO_FALTASFOLHA = {
    "Tipo": "Tipo",
    "Situacao": "Situacao",
    "Municipio": "Municipio",
    "Lotacao": "Lotacao",
    "Servidor": "Servidor",
    "CPF": "CPF",
    "TipoServidor": "TipoServidor",
    "Func": "Matricula",
    "Vinc": "Vinculo",
    "DataInicial": "Data Inicial",
    "DataFinal": "Data Final",
    "VigenteFunc": "FuncVigente",
    "VigenteVinc": "VincVigente",
    "ENCONTRADO_NO_SEAP": "EncontradoSEAP"
}

DICIONARIO_FALTASSEAP = {
    "Tipo": "Tipo",
    "Situacao": "Situacao",
    "Municipio": "Municipio",
    "Lotacao": "Lotacao",
    "Servidor": "Servidor",
    "CPF": "CPF",
    "TipoServidor": "TipoServidor",
    "Func": "Matricula",
    "Vinc": "Vinculo",
    "DataInicial": "Data Inicial",
    "DataFinal": "Data Final",
    "VigenteFunc": "FuncVigente",
    "VigenteVinc": "VincVigente",
    "ENCONTRADO_NO_SEAP": "EncontradoSEAP"
}

# ============================================================
# 2. DEFINIÇÕES DE CORES (HEXADECIMAL)
# ============================================================
COR_CABEÇALHO_VERDE = "004C9900"    # Verde escuro
COR_ZEBRA_VERDE_CLARO = "00E6F2E6"  # Verde claro
COR_BRANCA = "00FFFFFF"             # Branco

COR_CABEÇALHO_AZUL = "00003366"     # Azul escuro
COR_ZEBRA_AZUL_CLARO = "00E6F2FF"   # Azul claro


# ============================================================
# 3. NAVEGAÇÃO E LOCALIZAÇÃO DINÂMICA
# ============================================================
def obter_diretorio_raiz():
    """Sai da pasta 'processo' e retorna a raiz do projeto."""
    diretorio_processo = Path(__file__).resolve().parent
    return diretorio_processo.parent

def localizar_arquivo(nome_arquivo):
    """Procura o arquivo na raiz ou em suas subpastas."""
    raiz = obter_diretorio_raiz()
    
    # 1. Busca direta na raiz
    caminho_direto = raiz / nome_arquivo
    if caminho_direto.exists():
        return caminho_direto

    # 2. Busca recursiva caso esteja em subpastas
    arquivos_encontrados = list(raiz.rglob(nome_arquivo))
    if arquivos_encontrados:
        return arquivos_encontrados[0]
        
    return None


# ============================================================
# 4. APLICAÇÃO DE ESTILOS VISUAIS
# ============================================================
def aplicar_estilo_excel(caminho_arquivo, cor_cabecalho, cor_zebra):
    """Aplica cores, alinhamentos e largura de colunas usando openpyxl."""
    wb = load_workbook(caminho_arquivo)
    ws = wb.active

    fill_cabecalho = PatternFill(start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid")
    fonte_cabecalho = Font(color="FFFFFFFF", bold=True)
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    
    fill_zebra = PatternFill(start_color=cor_zebra, end_color=cor_zebra, fill_type="solid")
    fill_branco = PatternFill(start_color=COR_BRANCA, end_color=COR_BRANCA, fill_type="solid")

    # Formatar Cabeçalho (Linha 1)
    for cell in ws[1]:
        cell.fill = fill_cabecalho
        cell.font = fonte_cabecalho
        cell.alignment = alinhamento_centro

    # Formatar Linhas (Efeito Zebra)
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        linha_fill = fill_branco if row_idx % 2 == 0 else fill_zebra
        for cell in row:
            cell.fill = linha_fill
            cell.alignment = alinhamento_centro

    # Ajustar largura automática das colunas
    for col in ws.columns:
        max_length = 0
        coluna_letra = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[coluna_letra].width = max(max_length + 4, 12)

    wb.save(caminho_arquivo)


# ============================================================
# 5. EXECUÇÃO DA FORMATAÇÃO FOLHA (VERDE)
# ============================================================
def formatar_faltas_folha():
    print("Iniciando formatação da planilha FaltasFolha.xlsx (Tema Verde)...")
    
    caminho_arquivo = localizar_arquivo("FaltasFolha.xlsx")

    if not caminho_arquivo or not caminho_arquivo.exists():
        print("  -> ERRO: O arquivo FaltasFolha.xlsx não foi encontrado a partir do diretório raiz.")
        return

    print(f"  -> Arquivo encontrado em: {caminho_arquivo}")

    df = pd.read_excel(caminho_arquivo)
    df.columns = df.columns.str.strip()

    # Filtra mantendo a ordem exata das colunas do dicionário
    colunas_presentes = [col for col in DICIONARIO_FALTASFOLHA.keys() if col in df.columns]
    df = df[colunas_presentes]
    df = df.rename(columns=DICIONARIO_FALTASFOLHA)

    df.to_excel(caminho_arquivo, index=False)

    aplicar_estilo_excel(
        caminho_arquivo=caminho_arquivo,
        cor_cabecalho=COR_CABEÇALHO_VERDE,
        cor_zebra=COR_ZEBRA_VERDE_CLARO
    )
    print("  -> Formatação da Folha concluída com sucesso!\n")


# ============================================================
# 6. EXECUÇÃO DA FORMATAÇÃO SEAP (AZUL)
# ============================================================
def formatar_faltas_seap():
    print("Iniciando formatação da planilha FaltasSEAP.xlsx (Tema Azul)...")
    
    caminho_arquivo = localizar_arquivo("FaltasSEAP.xlsx")

    if not caminho_arquivo or not caminho_arquivo.exists():
        print("  -> ERRO: O arquivo FaltasSEAP.xlsx não foi encontrado a partir do diretório raiz.")
        return

    print(f"  -> Arquivo encontrado em: {caminho_arquivo}")

    df = pd.read_excel(caminho_arquivo)
    df.columns = df.columns.str.strip()

    # Filtra mantendo a ordem exata das colunas do dicionário
    colunas_presentes = [col for col in DICIONARIO_FALTASSEAP.keys() if col in df.columns]
    df = df[colunas_presentes]
    df = df.rename(columns=DICIONARIO_FALTASSEAP)

    df.to_excel(caminho_arquivo, index=False)

    aplicar_estilo_excel(
        caminho_arquivo=caminho_arquivo,
        cor_cabecalho=COR_CABEÇALHO_AZUL,
        cor_zebra=COR_ZEBRA_AZUL_CLARO
    )
    print("  -> Formatação do SEAP concluída com sucesso!\n")


if __name__ == "__main__":
    formatar_faltas_folha()
    formatar_faltas_seap()